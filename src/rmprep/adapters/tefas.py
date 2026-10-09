"""TEFAS adapter: reads files that a person exported by hand from the TEFAS site.

NO scraping and NO HTTP: this module never imports an HTTP client. Files in the
inbox (.xlsx, .csv) are copied byte-for-byte into the snapshot and parsed from
there. Headers are matched through `column_aliases` in sources.yaml; Turkish
number ("1.234,56") and date ("09.10.2026") formats are accepted.
"""

from __future__ import annotations

import io
import warnings
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from rmprep.adapters._inbox import list_inbox, payload_name
from rmprep.config import Config, TefasSource, normalize_header
from rmprep.contracts import TR_TZ, DataPoint
from rmprep.parsing import ParseError, detect_decimal, parse_date, parse_number, read_csv_rows
from rmprep.snapshot import Snapshot, SnapshotWriter, SourceRecord, raw_ref

_HEADER_SCAN_ROWS = 20


@dataclass(frozen=True)
class TefasRow:
    locator: str                 # e.g. "row=5" (csv) or "Sheet1!row=5" (xlsx)
    as_of: date
    entity: str                  # fund code, upper case
    values: dict[str, float]     # canonical value column -> value (empty cells left out)


def _xlsx_rows(data: bytes) -> tuple[str, list[list[Any]]]:
    from openpyxl import load_workbook

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")  # "Workbook contains no default style"
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # openpyxl raises many types for bad files
            raise ParseError(f"cannot open xlsx ({type(e).__name__})") from None
    try:
        ws = wb.worksheets[0]
        return ws.title, [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _find_header(rows: list[list[Any]], src: TefasSource) -> tuple[int, dict[str, int]]:
    lookup = {normalize_header(a): canon
              for canon, aliases in src.column_aliases.items() for a in aliases}
    for idx, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        colmap: dict[str, int] = {}
        for i, cell in enumerate(row):
            if cell is None:
                continue
            canon = lookup.get(normalize_header(str(cell)))
            if canon and canon not in colmap:
                colmap[canon] = i
        if src.date_column in colmap and src.entity_column in colmap:
            if not any(c in colmap for c in src.value_columns):
                raise ParseError("header row has no value column "
                                 f"(expected one of {list(src.value_columns)})")
            return idx, colmap
    raise ParseError("header row not found (need a date and a fund code column)")


def _blank(cell: Any) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


def parse_file(data: bytes, name: str, src: TefasSource) -> list[TefasRow]:
    ext = Path(name).suffix.lower()
    if ext == ".xlsx":
        sheet, rows = _xlsx_rows(data)
        prefix = f"{sheet}!"
    elif ext == ".csv":
        rows, prefix = read_csv_rows(data), ""
    else:
        raise ParseError(f"{name}: unsupported file type {ext!r} (use .xlsx or .csv)")

    try:
        header_idx, colmap = _find_header(rows, src)
    except ParseError as e:
        raise ParseError(f"{name}: {e}") from None

    def cell(row: list[Any], canon: str) -> Any:
        i = colmap[canon]
        return row[i] if i < len(row) else None

    value_cols = [c for c in src.value_columns if c in colmap]
    body = rows[header_idx + 1:]
    decimal = detect_decimal(
        cell(r, c) for r in body for c in value_cols if isinstance(cell(r, c), str))

    out: list[TefasRow] = []
    for n, row in enumerate(body, start=header_idx + 2):   # 1-based row number in the file
        if all(_blank(c) for c in row):
            continue
        loc = f"{prefix}row={n}"
        try:
            as_of = parse_date(cell(row, src.date_column))
            entity = str(cell(row, src.entity_column) or "").strip().upper()
            if not entity or "|" in entity:
                raise ParseError(f"bad fund code {entity!r}")
            values = {}
            for c in value_cols:
                v = parse_number(cell(row, c), decimal)
                if v is not None:
                    values[c] = v
        except ParseError as e:
            raise ParseError(f"{name} {loc}: {e}") from None
        out.append(TefasRow(locator=loc, as_of=as_of, entity=entity, values=values))
    return out


def fetch(
    src: TefasSource,
    writer: SnapshotWriter,
    *,
    inbox: Path,
    now: Callable[[], datetime] = lambda: datetime.now(TR_TZ),
) -> SourceRecord:
    """Copy all inbox files into the snapshot. All-or-nothing: one file that does
    not parse fails the source, and nothing is stored."""
    files = list_inbox(inbox, src.file_patterns)
    request = {"inbox": src.inbox, "files": [f.name for f in files],
               "method": "files exported by hand; no HTTP"}

    def failed(error: str) -> SourceRecord:
        return SourceRecord(status="failed", on_failure=src.on_failure,
                            request=request, error=error)

    if not files:
        return failed(f"no files matching {list(src.file_patterns)} in {src.inbox}")
    received = []
    for f in files:
        data = f.read_bytes()
        try:
            parse_file(data, f.name, src)
        except ParseError as e:
            return failed(str(e))
        received.append((f.name, data))

    fetched_at = now()
    refs = tuple(writer.add_payload(src.name, payload_name(name), data,
                                    meta={"original_name": name})
                 for name, data in received)
    return SourceRecord(status="ok", on_failure=src.on_failure, fetched_at=fetched_at,
                        request=request, payloads=refs)


def parse(cfg: Config, src: TefasSource, snap: Snapshot, *,
          out_snapshot_id: str | None = None, out_date: date | None = None) -> list[DataPoint]:
    rec = snap.sources.get(src.name)
    if rec is None or rec.status != "ok":
        raise ValueError(f"{snap.snapshot_id}: source {src.name!r} has no ok record")
    sid = out_snapshot_id or snap.snapshot_id
    points = []
    for p in rec.payloads:
        name = p.path.rsplit("/", 1)[-1]
        for row in parse_file(snap.read_payload(p.path), name, src):
            for canon, value in row.values.items():
                quantity = src.value_columns[canon]
                points.append(DataPoint.create(
                    quantity=quantity, entity=row.entity, value=value,
                    unit=cfg.quantities[quantity].unit, as_of=row.as_of, source=src.name,
                    source_url=src.source_url, fetched_at=rec.fetched_at, snapshot_id=sid,
                    raw_ref=raw_ref(snap.snapshot_id, p.path, f"{row.locator};col={canon}"),
                ))
    return points
